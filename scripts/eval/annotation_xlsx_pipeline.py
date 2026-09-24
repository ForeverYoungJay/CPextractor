import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

sys.path.append(str(Path(__file__).resolve().parents[2]))

from scripts.eval.export_usable_parameter_annotation import CSV_FIELDS as USABLE_PARAMETER_FIELDS
from scripts.eval.export_usable_parameter_annotation import build_usable_parameter_rows, _iter_paper_dirs


STATUS_VALUES = [
    "pending",
    "needs_adjudication",
    "correct",
    "wrong_material_object",
    "wrong_cp_model",
    "wrong_parameter_body",
    "wrong_scope",
    "wrong_evidence",
    "wrong_provenance",
    "insufficient_context",
    "spurious_record",
    "missing_record",
    "incorrect",
]

ERROR_TAG_VALUES = [
    "wrong_material_name",
    "wrong_constituent",
    "wrong_process_state",
    "wrong_model_family",
    "wrong_hardening_law",
    "wrong_parameter_name",
    "wrong_symbol",
    "wrong_value",
    "wrong_unit",
    "missing_unit_allowed",
    "wrong_scope_level",
    "wrong_scope_target",
    "wrong_condition",
    "wrong_temperature",
    "wrong_strain_rate",
    "wrong_table",
    "wrong_table_cell",
    "wrong_text_evidence",
    "unsupported_by_source",
    "wrong_origin_type",
    "wrong_reference",
    "wrong_calibration_context",
    "missing_required_context",
    "duplicate_record",
]

TRACKING_FIELDS = ["doi", "paper_dir", "claim_id", "record_index"]
ANNOTATION_FIELDS = ["annotation.status", "annotation.error_tags", "annotation.notes"]
LIST_FIELDS = {
    "annotation.error_tags",
    "parameter_scope.system_ids",
    "provenance.reference_ids",
    "provenance.adopted_from_reference_ids",
    "provenance.calibration_based_on_reference_ids",
}
BOOL_FIELDS = {"prediction_context.review_required"}
INT_FIELDS = {"record_index"}
FLOAT_FIELDS = {"parameter_body.value", "value", "value_SI", "evidence.page"}


def _require_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ModuleNotFoundError as exc:
        raise SystemExit(
            "openpyxl is required for XLSX annotation packets. "
            "Install it with `python3 -m pip install openpyxl`, or run this script with the bundled Codex Python."
        ) from exc
    return Workbook, load_workbook, Comment, Alignment, Font, PatternFill, get_column_letter, DataValidation


def _safe_id(value: str) -> str:
    return value.replace("/", "_").replace(":", "_")


def _load_jsonl(path: Path) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            payload = json.loads(line)
            if isinstance(payload, dict):
                records.append(payload)
    return records


def _load_json_records(path: Path) -> List[Dict[str, Any]]:
    if path.suffix.lower() == ".jsonl":
        return _load_jsonl(path)
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return [row for row in payload if isinstance(row, dict)]
    if isinstance(payload, dict):
        for key in ("records", "rows", "data"):
            value = payload.get(key)
            if isinstance(value, list):
                return [row for row in value if isinstance(row, dict)]
        return [payload]
    return []


def _write_jsonl(path: Path, records: Iterable[Dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for record in records:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")


def _get_nested(record: Dict[str, Any], dotted: str) -> Any:
    value: Any = record
    for part in dotted.split("."):
        if not isinstance(value, dict):
            return None
        value = value.get(part)
    return value


def _set_nested(record: Dict[str, Any], dotted: str, value: Any) -> None:
    parts = dotted.split(".")
    cur = record
    for part in parts[:-1]:
        node = cur.get(part)
        if not isinstance(node, dict):
            node = {}
            cur[part] = node
        cur = node
    cur[parts[-1]] = value


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "|".join(str(item) for item in value if item not in (None, ""))
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, tuple)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _parse_value(field: str, raw: Any) -> Any:
    if raw is None:
        return None
    text = str(raw).strip()
    if text == "":
        if field in LIST_FIELDS:
            return []
        return None
    if field in LIST_FIELDS:
        try:
            value = json.loads(text)
            if isinstance(value, list):
                return [str(item).strip() for item in value if str(item).strip()]
        except Exception:
            pass
        return [part.strip() for part in text.split("|") if part.strip()]
    if field in BOOL_FIELDS:
        return text.lower() in {"true", "1", "yes"}
    if field in INT_FIELDS:
        try:
            return int(text)
        except Exception:
            return text
    if field in FLOAT_FIELDS:
        try:
            return float(text)
        except Exception:
            return text
    return text


def _dedupe_fields(fields: Iterable[str]) -> List[str]:
    seen = set()
    out: List[str] = []
    for field in fields:
        if field in seen:
            continue
        seen.add(field)
        out.append(field)
    return out


def _default_fields(records: List[Dict[str, Any]]) -> List[str]:
    if records and isinstance(records[0].get("parameter_body"), dict):
        return _dedupe_fields([*TRACKING_FIELDS, *USABLE_PARAMETER_FIELDS])
    return _dedupe_fields(
        [
            "doi",
            "paper_dir",
            "claim_id",
            "record_index",
            "material_name",
            "canonical_name",
            "symbol",
            "domain",
            "value",
            "unit",
            "scope.scope",
            "scope.family_id",
            "evidence.file",
            "evidence.row_name",
            "evidence.column_name",
            "evidence.value_text",
            "annotation.status",
            "annotation.error_tags",
            "annotation.notes",
        ]
    )


def _write_taxonomy_sheet(wb: Any) -> Any:
    ws = wb.create_sheet("error_taxonomy")
    ws.append(["annotation.status", "annotation.error_tags"])
    max_len = max(len(STATUS_VALUES), len(ERROR_TAG_VALUES))
    for idx in range(max_len):
        ws.append([
            STATUS_VALUES[idx] if idx < len(STATUS_VALUES) else "",
            ERROR_TAG_VALUES[idx] if idx < len(ERROR_TAG_VALUES) else "",
        ])
    ws.sheet_state = "hidden"
    return ws


def _write_record_json_sheet(wb: Any, records: List[Dict[str, Any]]) -> None:
    ws = wb.create_sheet("_record_json")
    ws.append(["xlsx_row", "original_record_json"])
    for idx, record in enumerate(records, start=2):
        ws.append([idx, json.dumps(record, ensure_ascii=False)])
    ws.sheet_state = "veryHidden"


def export_records_to_xlsx(records: List[Dict[str, Any]], output_path: Path, fields: List[str] | None = None) -> None:
    Workbook, _, Comment, Alignment, Font, PatternFill, get_column_letter, DataValidation = _require_openpyxl()
    fields = fields or _default_fields(records)

    wb = Workbook()
    ws = wb.active
    ws.title = "annotations"
    taxonomy_ws = _write_taxonomy_sheet(wb)
    _write_record_json_sheet(wb, records)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    readonly_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append(fields)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = editable_fill if cell.value in ANNOTATION_FIELDS else header_fill
        if cell.value in TRACKING_FIELDS:
            cell.fill = readonly_fill
        if cell.value == "annotation.status":
            cell.comment = Comment("Choose the primary error class. Change pending to correct only after checking the source.", "CPextractor")
        if cell.value == "annotation.error_tags":
            cell.comment = Comment("Optional detail tags. For multiple tags, separate with |.", "CPextractor")

    for record in records:
        ws.append([_stringify(_get_nested(record, field)) for field in fields])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx, field in enumerate(fields, start=1):
        letter = get_column_letter(col_idx)
        sample_values = [field]
        for row_idx in range(2, min(len(records) + 2, 32)):
            sample_values.append(str(ws.cell(row=row_idx, column=col_idx).value or ""))
        width = max(10, min(48, max(len(value) for value in sample_values) + 2))
        if field in {"evidence.snippet", "parameter_scope.scope_target", "cp_model.model_label", "annotation.notes"}:
            width = min(72, max(width, 28))
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[1].height = 36

    for row in ws.iter_rows(min_row=2, max_row=len(records) + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    status_col = fields.index("annotation.status") + 1 if "annotation.status" in fields else None
    tags_col = fields.index("annotation.error_tags") + 1 if "annotation.error_tags" in fields else None
    max_row = max(len(records) + 1, 2000)
    if status_col:
        status_range = f"{taxonomy_ws.title}!$A$2:$A${len(STATUS_VALUES) + 1}"
        dv = DataValidation(type="list", formula1=f"={status_range}", allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(status_col)}2:{get_column_letter(status_col)}{max_row}")
    if tags_col:
        tag_range = f"{taxonomy_ws.title}!$B$2:$B${len(ERROR_TAG_VALUES) + 1}"
        dv = DataValidation(type="list", formula1=f"={tag_range}", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(tags_col)}2:{get_column_letter(tags_col)}{max_row}")

    instructions = wb.create_sheet("instructions")
    instructions.append(["Expert annotation workflow"])
    instructions.append(["1. Work in the annotations sheet."])
    instructions.append(["2. Use annotation.status for the primary error class."])
    instructions.append(["3. Use annotation.error_tags for optional detail tags. Excel dropdowns select one tag at a time; type tag1|tag2 for multiple tags."])
    instructions.append(["4. Use annotation.notes only when the correction or rationale is not obvious."])
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    instructions.column_dimensions["A"].width = 110

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _records_from_workbook(path: Path) -> Tuple[List[Dict[str, Any]], List[str]]:
    _, load_workbook, *_ = _require_openpyxl()
    wb = load_workbook(path)
    if "annotations" not in wb.sheetnames:
        raise ValueError(f"{path} does not contain an annotations sheet")
    ws = wb["annotations"]
    fields = [str(cell.value or "").strip() for cell in ws[1]]
    if not any(fields):
        return [], []

    original_by_row: Dict[int, Dict[str, Any]] = {}
    if "_record_json" in wb.sheetnames:
        meta = wb["_record_json"]
        for xlsx_row, payload in meta.iter_rows(min_row=2, values_only=True):
            if not xlsx_row or not payload:
                continue
            try:
                original_by_row[int(xlsx_row)] = json.loads(str(payload))
            except Exception:
                continue

    records: List[Dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        values = [ws.cell(row=row_idx, column=col_idx).value for col_idx in range(1, len(fields) + 1)]
        if not any(value not in (None, "") for value in values):
            continue
        record = original_by_row.get(row_idx, {})
        record = json.loads(json.dumps(record, ensure_ascii=False))
        for field, raw in zip(fields, values):
            if not field:
                continue
            _set_nested(record, field, _parse_value(field, raw))
        records.append(record)
    return records, fields


def export_from_input_root(input_root: Path, output_root: Path, *, per_paper: bool) -> Dict[str, int]:
    all_records: List[Dict[str, Any]] = []
    papers = 0
    for paper_dir in _iter_paper_dirs(input_root):
        records = build_usable_parameter_rows(paper_dir)
        if not records:
            continue
        papers += 1
        all_records.extend(records)
        if per_paper:
            doi = str(records[0].get("doi") or paper_dir.name.replace("_", "/"))
            export_records_to_xlsx(records, output_root / "packets" / _safe_id(doi) / "usable_parameters.xlsx")
    export_records_to_xlsx(all_records, output_root / "usable_parameter_annotation.xlsx")
    return {"papers": papers, "records": len(all_records)}


def export_from_jsonl(input_jsonl: Path, output_xlsx: Path) -> Dict[str, int]:
    records = _load_json_records(input_jsonl)
    export_records_to_xlsx(records, output_xlsx)
    return {"papers": 0, "records": len(records)}


def import_xlsx(input_xlsx: Path, output_jsonl: Path) -> Dict[str, int]:
    records, _ = _records_from_workbook(input_xlsx)
    _write_jsonl(output_jsonl, records)
    return {"records": len(records)}


def main() -> None:
    ap = argparse.ArgumentParser(description="Create dropdown-enabled XLSX annotation sheets from JSON/JSONL and import edited XLSX back to JSONL.")
    sub = ap.add_subparsers(dest="cmd", required=True)

    export_root = sub.add_parser("export-root", help="Build usable-parameter XLSX annotation workbooks from a fulltext input root.")
    export_root.add_argument("--input-root", required=True)
    export_root.add_argument("--output-root", required=True)
    export_root.add_argument("--per-paper", action="store_true", help="Also write one XLSX packet per paper.")

    export_jsonl = sub.add_parser(
        "export-json",
        aliases=["export-jsonl"],
        help="Build one XLSX annotation workbook from an annotation JSON or JSONL file.",
    )
    export_jsonl.add_argument("--input-json", default="", help="Annotation JSON or JSONL input path.")
    export_jsonl.add_argument("--input-jsonl", default="", help="Backward-compatible alias for --input-json.")
    export_jsonl.add_argument("--output-xlsx", required=True)

    import_cmd = sub.add_parser("import-xlsx", help="Import edited XLSX annotation workbook back to JSONL.")
    import_cmd.add_argument("--input-xlsx", required=True)
    import_cmd.add_argument("--output-jsonl", required=True)

    args = ap.parse_args()
    if args.cmd == "export-root":
        stats = export_from_input_root(Path(args.input_root), Path(args.output_root), per_paper=args.per_paper)
        print(f"Exported {stats['records']} records across {stats['papers']} papers -> {args.output_root}")
    elif args.cmd in {"export-json", "export-jsonl"}:
        input_json = args.input_json or args.input_jsonl
        if not input_json:
            raise SystemExit("export-json requires --input-json or --input-jsonl")
        stats = export_from_jsonl(Path(input_json), Path(args.output_xlsx))
        print(f"Exported {stats['records']} records -> {args.output_xlsx}")
    elif args.cmd == "import-xlsx":
        stats = import_xlsx(Path(args.input_xlsx), Path(args.output_jsonl))
        print(f"Imported {stats['records']} records -> {args.output_jsonl}")


if __name__ == "__main__":
    main()
