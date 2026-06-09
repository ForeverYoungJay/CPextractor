from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE_COLOR = RGBColor(0x00, 0x00, 0x00)
BODY_COLOR = RGBColor(0x00, 0x00, 0x00)
MUTED_COLOR = RGBColor(0x55, 0x55, 0x55)


def set_font(run, name: str, size_pt: float, color: RGBColor, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.15

    for style_name, size, before, after, color in [
        ("Heading 1", 20, 20, 6, TITLE_COLOR),
        ("Heading 2", 16, 18, 6, TITLE_COLOR),
        ("Heading 3", 14, 16, 4, RGBColor(0x43, 0x43, 0x43)),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = color
        pf = style.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.15

    if "Meta Custom" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("Meta Custom", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["Meta Custom"]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(10.5)
    style.font.color.rgb = MUTED_COLOR
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.15
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER


def parse_source(path: Path) -> tuple[dict[str, str], list[tuple[str, str]]]:
    metadata: dict[str, str] = {}
    blocks: list[tuple[str, str]] = []
    lines = path.read_text(encoding="utf-8").splitlines()

    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        if not line:
            idx += 1
            break
        if ":" in line:
            key, value = line.split(":", 1)
            metadata[key.strip()] = value.strip()
        idx += 1

    current: list[str] = []
    current_kind = "paragraph"

    def flush() -> None:
        nonlocal current, current_kind
        text = "\n".join(current).strip()
        if text:
            blocks.append((current_kind, text))
        current = []
        current_kind = "paragraph"

    for raw in lines[idx:]:
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("# "):
            flush()
            blocks.append(("h1", stripped[2:].strip()))
            continue
        if stripped.startswith("## "):
            flush()
            blocks.append(("h2", stripped[3:].strip()))
            continue
        if stripped.startswith("### "):
            flush()
            blocks.append(("h3", stripped[4:].strip()))
            continue
        if current:
            current.append(stripped)
        else:
            current = [stripped]
            current_kind = "paragraph"

    flush()
    return metadata, blocks


def add_title_block(doc: Document, metadata: dict[str, str]) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.line_spacing = 1.15
    title_run = title.add_run(metadata["Title"])
    set_font(title_run, "Arial", 26, TITLE_COLOR, bold=False)

    subtitle = doc.add_paragraph(style="Meta Custom")
    subtitle_run = subtitle.add_run(metadata["Subtitle"])
    set_font(subtitle_run, "Arial", 11, MUTED_COLOR)

    authors = doc.add_paragraph(style="Meta Custom")
    authors_run = authors.add_run(metadata["Authors"])
    set_font(authors_run, "Arial", 10.5, BODY_COLOR)

    affiliations = doc.add_paragraph(style="Meta Custom")
    affiliations_run = affiliations.add_run(metadata["Affiliations"])
    set_font(affiliations_run, "Arial", 10.5, MUTED_COLOR)

    correspondence = doc.add_paragraph(style="Meta Custom")
    correspondence.paragraph_format.space_after = Pt(10)
    correspondence_run = correspondence.add_run(metadata["Correspondence"])
    set_font(correspondence_run, "Arial", 10.5, MUTED_COLOR)


def build_doc(source: Path, output: Path) -> None:
    metadata, blocks = parse_source(source)
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    add_title_block(doc, metadata)

    for kind, text in blocks:
        if kind == "h1":
            doc.add_paragraph(text, style="Heading 1")
            continue
        if kind == "h2":
            doc.add_paragraph(text, style="Heading 2")
            continue
        if kind == "h3":
            doc.add_paragraph(text, style="Heading 3")
            continue

        para = doc.add_paragraph(style="Normal")
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.first_line_indent = Inches(0.0)
        run = para.add_run(text.replace("\n", " "))
        set_font(run, "Arial", 11, BODY_COLOR)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the CPextractor rewritten manuscript DOCX.")
    parser.add_argument("--source", default="manuscript/CPextractor_submission_rewrite.md")
    parser.add_argument("--output", default="output/documents/CPextractor_submission_rewrite.docx")
    args = parser.parse_args()

    build_doc(Path(args.source), Path(args.output))
    print(f"Saved DOCX -> {args.output}")


if __name__ == "__main__":
    main()
