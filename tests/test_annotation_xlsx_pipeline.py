import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

try:
    import openpyxl  # noqa: F401
except ModuleNotFoundError:
    openpyxl = None

from scripts.eval.annotation_xlsx_pipeline import export_from_jsonl, import_xlsx


@unittest.skipIf(openpyxl is None, "openpyxl is not installed")
class AnnotationXlsxPipelineTests(unittest.TestCase):
    def test_export_and_import_preserves_nested_annotation_edits(self):
        with TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_jsonl = root / "draft.jsonl"
            input_record = {
                "doi": "10.1000/example",
                "claim_id": "claim_1",
                "record_index": 0,
                "parameter_body": {
                    "canonical_name": "tau0",
                    "symbol": "tau0",
                    "value": 85,
                    "unit": "MPa",
                },
                "parameter_scope": {"scope_target": "family:basal"},
                "evidence": {"snippet": "tau0 85 MPa"},
                "provenance": {"origin_type": "calibrated"},
                "annotation": {"status": "correct", "error_tags": [], "notes": ""},
            }
            input_jsonl.write_text(json.dumps(input_record, ensure_ascii=False) + "\n", encoding="utf-8")

            xlsx_path = root / "annotation.xlsx"
            stats = export_from_jsonl(input_jsonl, xlsx_path)
            self.assertEqual({"papers": 0, "records": 1}, stats)

            wb = openpyxl.load_workbook(xlsx_path)
            self.assertIn("annotations", wb.sheetnames)
            self.assertIn("error_taxonomy", wb.sheetnames)
            self.assertIn("_record_json", wb.sheetnames)
            self.assertGreaterEqual(len(wb["annotations"].data_validations.dataValidation), 2)

            ws = wb["annotations"]
            headers = [cell.value for cell in ws[1]]
            status_col = headers.index("annotation.status") + 1
            tags_col = headers.index("annotation.error_tags") + 1
            notes_col = headers.index("annotation.notes") + 1
            ws.cell(row=2, column=status_col).value = "wrong_parameter_body"
            ws.cell(row=2, column=tags_col).value = "wrong_value|wrong_unit"
            ws.cell(row=2, column=notes_col).value = "Correct value is in Table 2."
            wb.save(xlsx_path)

            output_jsonl = root / "gold.jsonl"
            import_stats = import_xlsx(xlsx_path, output_jsonl)
            self.assertEqual({"records": 1}, import_stats)
            output_record = json.loads(output_jsonl.read_text(encoding="utf-8").strip())
            self.assertEqual("wrong_parameter_body", output_record["annotation"]["status"])
            self.assertEqual(["wrong_value", "wrong_unit"], output_record["annotation"]["error_tags"])
            self.assertEqual("Correct value is in Table 2.", output_record["annotation"]["notes"])
            self.assertEqual("tau0", output_record["parameter_body"]["canonical_name"])


if __name__ == "__main__":
    unittest.main()
